import json
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path
from tempfile import TemporaryDirectory

from argos import api, pitchbook
from argos.exporters import available_formats, export, render_deck_body, render_deck_html
from argos.server import create_server
from tests.support import seeded_connection


class HandlerTests(unittest.TestCase):
    """Handlers are plain functions, so most of the API needs no socket."""

    def setUp(self):
        self.conn = seeded_connection()
        self.addCleanup(self.conn.close)

    def test_health_reports_corpus_and_formats(self):
        payload = api.health(self.conn, {}, {})
        self.assertEqual(payload["status"], "ok")
        self.assertGreater(payload["documents"], 0)
        self.assertIn("html", payload["export_formats"])

    def test_search_parses_repeated_filters(self):
        params = {"q": ["revenue"], "doc_type": ["filing", "transcript"], "limit": ["5"]}
        payload = api.get_search(self.conn, params, {})
        self.assertLessEqual(len(payload["results"]), 5)
        self.assertTrue(all(row["doc_type"] in ("filing", "transcript") for row in payload["results"]))

    def test_document_not_found_raises_api_error(self):
        with self.assertRaises(api.ApiError) as caught:
            api.get_document(self.conn, {}, {}, 10 ** 7)
        self.assertEqual(caught.exception.status, 404)

    def test_company_profile_and_charts(self):
        profile = api.get_company(self.conn, {}, {}, "nvg")
        self.assertEqual(profile["ticker"], "NVG")
        charts = api.get_company_charts(self.conn, {}, {}, "NVG")
        self.assertEqual(set(charts), {"revenue", "margin", "sentiment", "volume"})
        self.assertTrue(all(value.startswith("<svg") or value.startswith("<p") for value in charts.values()))

    def test_ingest_requires_documents(self):
        with self.assertRaises(api.ApiError):
            api.post_ingest(self.conn, {}, {"documents": []})

    def test_ingest_accepts_documents(self):
        result = api.post_ingest(self.conn, {}, {"documents": [{
            "external_id": "api-1", "title": "Ad hoc note", "doc_type": "memo",
            "published_at": "2026-09-01", "body": "Margins expanded on pricing discipline.",
        }]})
        self.assertEqual(result["added"], 1)

    def test_watchlist_round_trip(self):
        created = api.post_watchlist(self.conn, {}, {"name": "Grid software", "tickers": ["NVG"]})
        api.post_watchlist_items(self.conn, {}, {"tickers": ["HLDX"]}, created["id"])
        listed = api.get_watchlists(self.conn, {}, {})["watchlists"][0]
        self.assertEqual({row["ticker"] for row in listed["companies"]}, {"NVG", "HLDX"})
        api.delete_watchlist_item(self.conn, {}, {}, created["id"], "HLDX")
        listed = api.get_watchlists(self.conn, {}, {})["watchlists"][0]
        self.assertEqual({row["ticker"] for row in listed["companies"]}, {"NVG"})
        api.delete_watchlist(self.conn, {}, {}, created["id"])
        self.assertEqual(api.get_watchlists(self.conn, {}, {})["watchlists"], [])

    def test_unknown_ticker_on_watchlist_is_rejected(self):
        created = api.post_watchlist(self.conn, {}, {"name": "L"})
        with self.assertRaises(api.ApiError):
            api.post_watchlist_items(self.conn, {}, {"tickers": ["NOPE"]}, created["id"])

    def test_saved_search_alerts_report_new_documents(self):
        saved = api.post_saved_search(self.conn, {}, {
            "name": "Pricing", "query": "pricing pressure", "last_seen_document_id": 0,
        })
        alerts = api.get_alerts(self.conn, {}, {})["alerts"]
        self.assertEqual(alerts[0]["saved_search_id"], saved["id"])
        self.assertGreater(alerts[0]["new_documents"], 0)
        # Once marked seen, the same documents stop alerting.
        api.get_alerts(self.conn, {"mark_seen": ["1"]}, {})
        self.assertEqual(api.get_alerts(self.conn, {}, {})["alerts"][0]["new_documents"], 0)

    def test_saved_search_requires_a_query(self):
        with self.assertRaises(api.ApiError):
            api.post_saved_search(self.conn, {}, {"name": "x"})

    def test_pitchbook_lifecycle(self):
        book = api.post_pitchbook(self.conn, {}, {"target_ticker": "NVG", "title": "Project API"})
        self.assertTrue(book["slides"])
        preview = api.get_pitchbook_preview(self.conn, {}, {}, book["id"])
        self.assertIn("<section", preview["html"])
        valuation = api.get_valuation(self.conn, {}, {}, book["id"])
        self.assertTrue(valuation["ranges"])
        charts = api.get_valuation_chart(self.conn, {}, {}, book["id"])
        self.assertIn("<svg", charts["football_field"])
        api.delete_pitchbook(self.conn, {}, {}, book["id"])
        with self.assertRaises(api.ApiError):
            api.get_pitchbook(self.conn, {}, {}, book["id"])

    def test_pitchbook_requires_a_target(self):
        with self.assertRaises(api.ApiError):
            api.post_pitchbook(self.conn, {}, {"title": "No target"})

    def test_unknown_target_returns_404(self):
        with self.assertRaises(api.ApiError) as caught:
            api.post_pitchbook(self.conn, {}, {"target_ticker": "NOPE"})
        self.assertEqual(caught.exception.status, 404)


class HttpTests(unittest.TestCase):
    """A live server, to cover routing, status codes and downloads."""

    @classmethod
    def setUpClass(cls):
        cls.conn = seeded_connection()
        cls.server = create_server("127.0.0.1", 0, conn=cls.conn)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.base = f"http://127.0.0.1:{cls.server.server_address[1]}"
        cls.book_id = pitchbook.create_pitchbook(cls.conn, title="Project HTTP", target_ticker="NVG")

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.conn.close()

    def get(self, path):
        with urllib.request.urlopen(self.base + path) as response:
            return response.status, response.read(), dict(response.headers)

    def post(self, path, payload):
        request = urllib.request.Request(
            self.base + path, data=json.dumps(payload).encode(),
            headers={"Content-Type": "application/json"}, method="POST",
        )
        with urllib.request.urlopen(request) as response:
            return response.status, json.loads(response.read())

    def test_index_and_assets_are_served(self):
        for path, marker in (("/", b"<title>Argos"), ("/app.js", b"const state"), ("/app.css", b"--surface-1")):
            status, body, _ = self.get(path)
            self.assertEqual(status, 200)
            self.assertIn(marker, body)

    def test_palette_css_defines_both_schemes(self):
        _, body, headers = self.get("/tokens.css")
        text = body.decode()
        self.assertIn("text/css", headers["Content-Type"])
        self.assertIn("prefers-color-scheme: dark", text)
        self.assertIn('[data-theme="dark"]', text)

    def test_static_traversal_is_blocked(self):
        with self.assertRaises(urllib.error.HTTPError) as caught:
            self.get("/../argos/config.py")
        self.assertEqual(caught.exception.code, 404)

    def test_search_endpoint(self):
        status, body, _ = self.get("/api/search?q=margin&limit=3")
        payload = json.loads(body)
        self.assertEqual(status, 200)
        self.assertLessEqual(len(payload["results"]), 3)

    def test_unknown_endpoint_returns_404(self):
        with self.assertRaises(urllib.error.HTTPError) as caught:
            self.get("/api/nope")
        self.assertEqual(caught.exception.code, 404)

    def test_method_not_allowed(self):
        with self.assertRaises(urllib.error.HTTPError) as caught:
            self.post("/api/overview", {})
        self.assertEqual(caught.exception.code, 405)
        self.assertIn("allowed", json.loads(caught.exception.read()))

    def test_invalid_json_body(self):
        request = urllib.request.Request(
            self.base + "/api/pitchbooks", data=b"{not json", method="POST",
            headers={"Content-Type": "application/json"},
        )
        with self.assertRaises(urllib.error.HTTPError) as caught:
            urllib.request.urlopen(request)
        self.assertEqual(caught.exception.code, 400)

    def test_export_download_sets_attachment_headers(self):
        status, body, headers = self.get(f"/api/pitchbooks/{self.book_id}/export?format=html")
        self.assertEqual(status, 200)
        self.assertIn("attachment;", headers["Content-Disposition"])
        self.assertIn(b"<!doctype html>", body[:200].lower())

    def test_export_metadata_without_download(self):
        _, body, _ = self.get(f"/api/pitchbooks/{self.book_id}/export?format=html&download=0")
        payload = json.loads(body)
        self.assertEqual(payload["format"], "html")
        self.assertGreater(payload["bytes"], 0)

    def test_unknown_export_format_is_rejected(self):
        with self.assertRaises(urllib.error.HTTPError) as caught:
            self.get(f"/api/pitchbooks/{self.book_id}/export?format=xlsx&download=0")
        self.assertEqual(caught.exception.code, 400)


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.conn = seeded_connection()
        book_id = pitchbook.create_pitchbook(
            cls.conn, title="Project Export", target_ticker="NVG", client_name="Board",
        )
        pitchbook.add_evidence(cls.conn, book_id, quote="Churn ticked up to 5.2%.", document_id=1)
        cls.book = pitchbook.get_pitchbook(cls.conn, book_id)

    @classmethod
    def tearDownClass(cls):
        cls.conn.close()

    def test_html_deck_is_self_contained(self):
        markup = render_deck_html(self.book)
        self.assertIn("<!doctype html>", markup.lower())
        self.assertEqual(markup.count("<section class=\"slide"), len(self.book["slides"]))
        self.assertNotIn("http://", markup.replace("http://www.w3.org", ""))  # no external assets
        self.assertIn("prefers-color-scheme: dark", markup)

    def test_deck_body_matches_slide_count(self):
        body = render_deck_body(self.book)
        self.assertEqual(body.count("<section class=\"slide"), len(self.book["slides"]))

    def test_html_escapes_untrusted_text(self):
        book = dict(self.book)
        book["slides"] = [{
            "kind": "process", "title": "<script>alert(1)</script>", "subtitle": None,
            "payload": {"bullets": ["<img src=x onerror=alert(1)>"]},
        }]
        markup = render_deck_body(book)
        self.assertNotIn("<script>alert(1)</script>", markup)
        self.assertNotIn("<img src=x", markup)

    def test_every_available_format_writes_a_file(self):
        formats = available_formats()
        with TemporaryDirectory() as tmp:
            for fmt, supported in formats.items():
                if not supported:
                    continue
                path = export(self.book, fmt, Path(tmp) / f"deck.{fmt}")
                self.assertTrue(path.exists(), fmt)
                self.assertGreater(path.stat().st_size, 1000, fmt)

    def test_unsupported_format_raises(self):
        with TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                export(self.book, "xlsx", Path(tmp) / "deck.xlsx")

    @unittest.skipUnless(available_formats()["pptx"], "python-pptx not installed")
    def test_pptx_has_one_slide_per_payload_with_native_charts(self):
        from pptx import Presentation

        from argos.exporters.pptx_export import build_presentation

        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            build_presentation(self.book).save(str(path))
            presentation = Presentation(str(path))
            self.assertEqual(len(presentation.slides._sldIdLst), len(self.book["slides"]))
            charts = sum(1 for slide in presentation.slides for shape in slide.shapes if shape.has_chart)
            tables = sum(1 for slide in presentation.slides for shape in slide.shapes if shape.has_table)
            self.assertGreaterEqual(charts, 4)
            self.assertGreaterEqual(tables, 5)


if __name__ == "__main__":
    unittest.main()
