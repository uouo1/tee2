"""PowerPoint export.

Charts are native PowerPoint charts, not pictures, so a banker can double-click
a comps bar and retype a number at 11pm without coming back to the application.
The football field uses the stacked-bar trick (an invisible base series carrying
the low value, a visible series carrying the span), which is how the chart is
built by hand anyway.

Requires ``python-pptx``; :func:`argos.exporters.available_formats` reports
whether it is installed.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Sequence

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .. import charts as chart_tokens

LIGHT = chart_tokens.PALETTE["light"]


def _rgb(hex_colour: str) -> RGBColor:
    return RGBColor.from_string(hex_colour.lstrip("#").upper())


SERIES_1 = _rgb(LIGHT["series_1"])
INK = _rgb(LIGHT["text"])
INK_SECONDARY = _rgb(LIGHT["text_secondary"])
MUTED = _rgb(LIGHT["muted"])
SEQ = [_rgb(step) for step in LIGHT["seq"]]

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.62)
CONTENT_TOP = Inches(1.45)
CONTENT_W = SLIDE_W - MARGIN * 2
CONTENT_H = SLIDE_H - CONTENT_TOP - Inches(0.5)

CURRENCY_SYMBOLS = {"USD": "$", "EUR": "€", "GBP": "£"}


def _symbol(currency: str | None) -> str:
    return CURRENCY_SYMBOLS.get((currency or "USD").upper(), "")


def _fmt(value: Any, pattern: str = "{:,.0f}") -> str:
    if value is None:
        return "n/a"
    try:
        return pattern.format(float(value))
    except (TypeError, ValueError):
        return str(value)


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _heading(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(MARGIN, Inches(0.42), CONTENT_W, Inches(0.9))
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.text = title
    para.font.size = Pt(26)
    para.font.bold = True
    para.font.color.rgb = INK
    if subtitle:
        sub = frame.add_paragraph()
        sub.text = subtitle
        sub.font.size = Pt(12.5)
        sub.font.color.rgb = INK_SECONDARY
    rule = slide.shapes.add_shape(1, MARGIN, Inches(1.33), CONTENT_W, Emu(9525))
    rule.fill.solid()
    rule.fill.fore_color.rgb = _rgb(LIGHT["axis"])
    rule.line.fill.background()


def _bullets(slide, bullets: Sequence[Any], *, top=CONTENT_TOP, height=CONTENT_H, width=CONTENT_W) -> None:
    box = slide.shapes.add_textbox(MARGIN, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    first = True
    for bullet in bullets:
        text = bullet if isinstance(bullet, str) else bullet.get("text", "")
        source = None if isinstance(bullet, str) else bullet.get("source")
        para = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        para.text = f"•  {text}"
        para.font.size = Pt(14)
        para.font.color.rgb = INK
        para.space_after = Pt(7)
        if source:
            cite = frame.add_paragraph()
            cite.text = (
                f"     Source: {source.get('title', '')} · {source.get('doc_type', '')} · "
                f"{source.get('published_at', '')} [doc {source.get('document_id')}]"
            )
            cite.font.size = Pt(9.5)
            cite.font.color.rgb = MUTED
            cite.space_after = Pt(8)


def _table(slide, columns: Sequence[dict[str, Any]], rows: Sequence[dict[str, Any]], *,
           top=CONTENT_TOP, height=None, width=CONTENT_W, left=MARGIN,
           highlight_key: str | None = None, font_size: float = 9.5) -> None:
    row_count = len(rows) + 1
    height = height or min(CONTENT_H, Inches(0.3) * row_count)
    shape = slide.shapes.add_table(row_count, len(columns), left, top, width, height)
    table = shape.table
    for index, column in enumerate(columns):
        cell = table.cell(0, index)
        cell.text = column["label"]
        para = cell.text_frame.paragraphs[0]
        para.font.size = Pt(font_size)
        para.font.bold = True
        para.font.color.rgb = INK_SECONDARY
        if column.get("numeric"):
            para.alignment = PP_ALIGN.RIGHT
    for row_index, row in enumerate(rows, start=1):
        highlighted = bool(highlight_key and row.get(highlight_key))
        for col_index, column in enumerate(columns):
            cell = table.cell(row_index, col_index)
            cell.text = _fmt(row.get(column["key"]), column["format"]) if column.get("format") else str(row.get(column["key"], "") or "")
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(font_size)
            para.font.bold = highlighted
            para.font.color.rgb = INK
            if column.get("numeric"):
                para.alignment = PP_ALIGN.RIGHT


def _note(slide, text: str, *, top) -> None:
    if not text:
        return
    box = slide.shapes.add_textbox(MARGIN, top, CONTENT_W, Inches(0.8))
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.text = text
    para.font.size = Pt(11)
    para.font.color.rgb = INK_SECONDARY


def _style_axes(chart) -> None:
    for axis in (chart.category_axis, chart.value_axis):
        axis.has_major_gridlines = axis is chart.value_axis
        axis.tick_labels.font.size = Pt(10)
        axis.tick_labels.font.color.rgb = MUTED
        axis.format.line.color.rgb = _rgb(LIGHT["axis"])


def _add_column_chart(slide, categories: Sequence[str], values: Sequence[float], *,
                      left=MARGIN, top=CONTENT_TOP, width=CONTENT_W, height=Inches(3.1),
                      series_name: str = "Value", number_format: str = "#,##0") -> None:
    data = CategoryChartData()
    data.categories = list(categories)
    data.add_series(series_name, tuple(float(v or 0) for v in values), number_format)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width, height, data)
    chart = frame.chart
    chart.has_legend = False
    chart.plots[0].gap_width = 60
    series = chart.plots[0].series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = SERIES_1
    _style_axes(chart)


def _add_line_chart(slide, categories: Sequence[str], values: Sequence[float], *,
                    left=MARGIN, top=CONTENT_TOP, width=CONTENT_W, height=Inches(3.0),
                    series_name: str = "Value", number_format: str = "0.0%") -> None:
    data = CategoryChartData()
    data.categories = list(categories)
    data.add_series(series_name, tuple(values), number_format)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, left, top, width, height, data)
    chart = frame.chart
    chart.has_legend = False
    series = chart.plots[0].series[0]
    series.format.line.color.rgb = SERIES_1
    series.format.line.width = Pt(2)
    _style_axes(chart)


def _add_football_field(slide, ranges: Sequence[dict[str, Any]], *, reference_price: float | None,
                        currency: str, top=CONTENT_TOP, height=Inches(3.5)) -> None:
    """Floating bars via a transparent base series plus a visible span series."""
    data = CategoryChartData()
    data.categories = [r["method"] for r in reversed(list(ranges))]
    ordered = list(reversed(list(ranges)))
    data.add_series("Base", tuple(float(r["low"]) for r in ordered), "#,##0.00")
    data.add_series("Range", tuple(float(r["high"]) - float(r["low"]) for r in ordered), "#,##0.00")
    frame = slide.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, MARGIN, top, CONTENT_W, height, data)
    chart = frame.chart
    chart.has_legend = False
    base, span = chart.plots[0].series
    base.format.fill.background()
    base.format.line.fill.background()
    span.format.fill.solid()
    span.format.fill.fore_color.rgb = SERIES_1
    _style_axes(chart)
    label_top = Emu(int(top) + int(height))
    symbol = _symbol(currency)
    detail = "; ".join(
        f"{r['method']}: {symbol}{float(r['low']):,.2f}–{symbol}{float(r['high']):,.2f}" for r in ranges
    )
    reference = f"Current share price: {symbol}{reference_price:,.2f}. " if reference_price else ""
    _note(slide, f"{reference}{detail}", top=label_top)


def _add_heatmap_table(slide, payload: dict[str, Any]) -> None:
    grid = payload.get("grid", [])
    wacc_axis = payload.get("wacc_axis", [])
    terminal_axis = payload.get("terminal_axis", [])
    flat = [float(v) for row in grid for v in row if v is not None]
    low, high = (min(flat), max(flat)) if flat else (0.0, 1.0)
    span = (high - low) or 1.0
    shape = slide.shapes.add_table(len(grid) + 1, len(wacc_axis) + 1, MARGIN, CONTENT_TOP, CONTENT_W, Inches(0.36) * (len(grid) + 1))
    table = shape.table
    corner = table.cell(0, 0)
    corner.text = "Terminal growth \\ WACC"
    corner.text_frame.paragraphs[0].font.size = Pt(9)
    corner.text_frame.paragraphs[0].font.color.rgb = INK_SECONDARY
    for index, wacc in enumerate(wacc_axis, start=1):
        cell = table.cell(0, index)
        cell.text = f"{float(wacc) * 100:.2f}%"
        para = cell.text_frame.paragraphs[0]
        para.font.size = Pt(10)
        para.font.bold = True
        para.font.color.rgb = INK_SECONDARY
        para.alignment = PP_ALIGN.CENTER
    for row_index, (growth, row) in enumerate(zip(terminal_axis, grid), start=1):
        label = table.cell(row_index, 0)
        label.text = f"{float(growth) * 100:.2f}%"
        label.text_frame.paragraphs[0].font.size = Pt(10)
        label.text_frame.paragraphs[0].font.color.rgb = INK_SECONDARY
        for col_index, value in enumerate(row, start=1):
            cell = table.cell(row_index, col_index)
            cell.text = "n/a" if value is None else f"{float(value):,.2f}"
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(10)
            para.alignment = PP_ALIGN.CENTER
            if value is None:
                continue
            step = min(6, max(0, int((float(value) - low) / span * 6)))
            cell.fill.solid()
            cell.fill.fore_color.rgb = SEQ[step]
            para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if step >= 3 else INK


# ------------------------------------------------------------- slide writers


def _slide_cover(prs, slide_data: dict[str, Any]) -> None:
    slide = _blank(prs)
    payload = slide_data.get("payload", {})
    box = slide.shapes.add_textbox(MARGIN, Inches(2.2), CONTENT_W, Inches(3.0))
    frame = box.text_frame
    frame.word_wrap = True
    badge = frame.paragraphs[0]
    badge.text = "CONFIDENTIAL — DISCUSSION MATERIALS"
    badge.font.size = Pt(11)
    badge.font.color.rgb = MUTED
    title = frame.add_paragraph()
    title.text = slide_data.get("title", "")
    title.font.size = Pt(40)
    title.font.bold = True
    title.font.color.rgb = INK
    if slide_data.get("subtitle"):
        sub = frame.add_paragraph()
        sub.text = slide_data["subtitle"]
        sub.font.size = Pt(16)
        sub.font.color.rgb = INK_SECONDARY
    meta = frame.add_paragraph()
    meta.text = " · ".join(
        str(part) for part in (
            f"Prepared for {payload.get('client_name')}" if payload.get("client_name") else None,
            payload.get("target"), payload.get("ticker"), payload.get("banker"), payload.get("date"),
        ) if part
    )
    meta.font.size = Pt(12)
    meta.font.color.rgb = MUTED


def _slide_text(prs, slide_data: dict[str, Any]) -> None:
    slide = _blank(prs)
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    payload = slide_data.get("payload", {})
    if payload.get("bullets"):
        _bullets(slide, payload["bullets"])
    else:
        _bullets(slide, [payload.get("body", "")])


def _slide_executive_summary(prs, slide_data: dict[str, Any]) -> None:
    slide = _blank(prs)
    payload = slide_data.get("payload", {})
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    reference = payload.get("reference_range") or {}
    symbol = _symbol(payload.get("currency"))
    if reference:
        box = slide.shapes.add_textbox(MARGIN, Inches(1.4), CONTENT_W, Inches(0.5))
        para = box.text_frame.paragraphs[0]
        para.text = (
            f"Reference range {symbol}{reference.get('low', 0):,.2f} – {symbol}{reference.get('high', 0):,.2f} "
            f"per share  ·  current {symbol}{payload.get('share_price', 0):,.2f}"
        )
        para.font.size = Pt(13)
        para.font.bold = True
        para.font.color.rgb = SERIES_1
    _bullets(slide, payload.get("bullets", []), top=Inches(2.0), height=Inches(4.9))


def _slide_company_overview(prs, slide_data: dict[str, Any]) -> None:
    slide = _blank(prs)
    payload = slide_data.get("payload", {})
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    box = slide.shapes.add_textbox(MARGIN, CONTENT_TOP, Inches(6.4), Inches(4.4))
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.text = payload.get("description", "")
    para.font.size = Pt(12.5)
    para.font.color.rgb = INK
    extracted = payload.get("extracted_kpis") or []
    if extracted:
        heading = frame.add_paragraph()
        heading.text = "\nMetrics extracted from the corpus"
        heading.font.size = Pt(11)
        heading.font.bold = True
        heading.font.color.rgb = INK_SECONDARY
        for kpi in extracted:
            line = frame.add_paragraph()
            line.text = f"•  {kpi['label']}: {kpi['value']} ({kpi['observations']} observations)"
            line.font.size = Pt(11)
            line.font.color.rgb = INK
    rows = [{"label": k["label"], "value": k["value"]} for k in payload.get("kpis", [])]
    shape = slide.shapes.add_table(len(rows), 2, Inches(7.3), CONTENT_TOP, Inches(5.4), Inches(0.33) * len(rows))
    for index, row in enumerate(rows):
        for col, key in enumerate(("label", "value")):
            cell = shape.table.cell(index, col)
            cell.text = str(row[key])
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(10.5)
            para.font.color.rgb = INK_SECONDARY if col == 0 else INK
            para.font.bold = col == 1
            if col == 1:
                para.alignment = PP_ALIGN.RIGHT


def _slide_market_overview(prs, slide_data: dict[str, Any]) -> None:
    payload = slide_data.get("payload", {})
    slide = _blank(prs)
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    box = slide.shapes.add_textbox(MARGIN, Inches(1.4), CONTENT_W, Inches(0.9))
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.text = payload.get("narrative", "")
    para.font.size = Pt(12)
    para.font.color.rgb = INK
    timeline = payload.get("sentiment_timeline") or []
    if timeline:
        _add_line_chart(
            slide, [row["month"] for row in timeline], [float(row["sentiment"] or 0) for row in timeline],
            left=MARGIN, top=Inches(2.35), width=Inches(6.0), height=Inches(2.9),
            series_name="Mean tone", number_format="0.00",
        )
    themes = payload.get("themes") or []
    if themes:
        _table(
            slide,
            [
                {"key": "theme", "label": "Emerging theme"},
                {"key": "mentions", "label": "Recent", "numeric": True, "format": "{:,.0f}"},
                {"key": "prior_mentions", "label": "Prior", "numeric": True, "format": "{:,.0f}"},
                {"key": "lift", "label": "Lift", "numeric": True, "format": "{:,.2f}x"},
            ],
            themes[:8], top=Inches(2.35), width=Inches(5.7), left=Inches(7.0),
        )


def _slide_financials(prs, slide_data: dict[str, Any]) -> None:
    payload = slide_data.get("payload", {})
    rows = payload.get("rows", [])
    currency = payload.get("currency", "USD")
    slide = _blank(prs)
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    categories = [f"FY{r['fiscal_year']}" for r in rows]
    _add_column_chart(
        slide, categories, [r["revenue"] for r in rows],
        left=MARGIN, top=CONTENT_TOP, width=Inches(6.0), height=Inches(2.9),
        series_name=f"Revenue ({currency}m)",
    )
    _add_line_chart(
        slide, categories, [r.get("ebitda_margin") for r in rows],
        left=Inches(7.0), top=CONTENT_TOP, width=Inches(5.7), height=Inches(2.9),
        series_name="EBITDA margin",
    )
    _table(
        slide,
        [
            {"key": "fiscal_year", "label": "FY", "format": "{:,.0f}"},
            {"key": "revenue", "label": "Revenue", "numeric": True, "format": "{:,.0f}"},
            {"key": "ebitda", "label": "EBITDA", "numeric": True, "format": "{:,.0f}"},
            {"key": "ebitda_margin", "label": "Margin", "numeric": True, "format": "{:.1%}"},
            {"key": "ebit", "label": "EBIT", "numeric": True, "format": "{:,.0f}"},
            {"key": "net_income", "label": "Net income", "numeric": True, "format": "{:,.0f}"},
            {"key": "eps", "label": "EPS", "numeric": True, "format": "{:,.2f}"},
            {"key": "capex", "label": "Capex", "numeric": True, "format": "{:,.0f}"},
        ],
        rows, top=Inches(4.6),
    )
    _note(slide, payload.get("note", ""), top=Inches(6.9))


def _slide_valuation_summary(prs, slide_data: dict[str, Any]) -> None:
    payload = slide_data.get("payload", {})
    slide = _blank(prs)
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    _add_football_field(
        slide, payload.get("ranges", []),
        reference_price=payload.get("reference_price"),
        currency=payload.get("currency", "USD"),
        top=CONTENT_TOP, height=Inches(3.9),
    )


def _slide_trading_comps(prs, slide_data: dict[str, Any]) -> None:
    payload = slide_data.get("payload", {})
    currency = payload.get("currency", "USD")
    target = dict(payload.get("target") or {})
    rows = ([target] if target else []) + [dict(r) for r in payload.get("rows", [])]
    slide = _blank(prs)
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    _table(
        slide,
        [
            {"key": "ticker", "label": "Ticker"},
            {"key": "name", "label": "Company"},
            {"key": "share_price", "label": "Price", "numeric": True, "format": "{:,.2f}"},
            {"key": "market_cap", "label": f"Mkt cap ({currency}m)", "numeric": True, "format": "{:,.0f}"},
            {"key": "enterprise_value", "label": f"EV ({currency}m)", "numeric": True, "format": "{:,.0f}"},
            {"key": "revenue", "label": "Revenue", "numeric": True, "format": "{:,.0f}"},
            {"key": "ebitda_margin", "label": "EBITDA margin", "numeric": True, "format": "{:.1%}"},
            {"key": "ev_revenue", "label": "EV/Rev", "numeric": True, "format": "{:,.1f}x"},
            {"key": "ev_ebitda", "label": "EV/EBITDA", "numeric": True, "format": "{:,.1f}x"},
            {"key": "pe", "label": "P/E", "numeric": True, "format": "{:,.1f}x"},
        ],
        rows, top=CONTENT_TOP, highlight_key="is_target",
    )
    stats = payload.get("stats", {})
    stat_rows = [
        {
            "multiple": label,
            **{key: stats.get(metric, {}).get(key) for key in ("min", "p25", "median", "p75", "max")},
        }
        for metric, label in (("ev_revenue", "EV / Revenue"), ("ev_ebitda", "EV / EBITDA"), ("pe", "P / E"))
    ]
    _table(
        slide,
        [
            {"key": "multiple", "label": "Peer distribution"},
            {"key": "min", "label": "Min", "numeric": True, "format": "{:,.1f}x"},
            {"key": "p25", "label": "25th", "numeric": True, "format": "{:,.1f}x"},
            {"key": "median", "label": "Median", "numeric": True, "format": "{:,.1f}x"},
            {"key": "p75", "label": "75th", "numeric": True, "format": "{:,.1f}x"},
            {"key": "max", "label": "Max", "numeric": True, "format": "{:,.1f}x"},
        ],
        stat_rows, top=Inches(5.35),
    )


def _slide_precedents(prs, slide_data: dict[str, Any]) -> None:
    payload = slide_data.get("payload", {})
    slide = _blank(prs)
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    _table(
        slide,
        [
            {"key": "announced_on", "label": "Announced"},
            {"key": "target", "label": "Target"},
            {"key": "acquirer", "label": "Acquirer"},
            {"key": "ev", "label": "EV (USDm)", "numeric": True, "format": "{:,.0f}"},
            {"key": "ev_revenue", "label": "EV/Rev", "numeric": True, "format": "{:,.1f}x"},
            {"key": "ev_ebitda", "label": "EV/EBITDA", "numeric": True, "format": "{:,.1f}x"},
            {"key": "premium_1day", "label": "Premium", "numeric": True, "format": "{:.0%}"},
            {"key": "consideration", "label": "Consideration"},
        ],
        payload.get("rows", []), top=CONTENT_TOP,
    )
    _note(slide, payload.get("note", ""), top=Inches(5.6))


def _slide_dcf(prs, slide_data: dict[str, Any]) -> None:
    payload = slide_data.get("payload", {})
    projections = payload.get("projections", [])
    currency = payload.get("currency", "USD")
    slide = _blank(prs)
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    _add_column_chart(
        slide, [str(p["year"]) for p in projections], [p["fcf"] for p in projections],
        left=MARGIN, top=CONTENT_TOP, width=Inches(5.9), height=Inches(2.8),
        series_name=f"Unlevered FCF ({currency}m)",
    )
    symbol = _symbol(currency)
    bridge = [
        ("PV of forecast FCF", f"{symbol}{payload.get('pv_explicit', 0):,.0f}m"),
        ("PV of terminal value", f"{symbol}{payload.get('pv_terminal', 0):,.0f}m"),
        ("Enterprise value", f"{symbol}{payload.get('enterprise_value', 0):,.0f}m"),
        ("Equity value", f"{symbol}{payload.get('equity_value', 0):,.0f}m"),
        ("Value per share", f"{symbol}{payload.get('value_per_share', 0):,.2f}"),
        ("Current share price", f"{symbol}{payload.get('share_price', 0):,.2f}"),
    ]
    shape = slide.shapes.add_table(len(bridge), 2, Inches(7.0), CONTENT_TOP, Inches(5.7), Inches(0.33) * len(bridge))
    for index, (label, value) in enumerate(bridge):
        for col, text in enumerate((label, value)):
            cell = shape.table.cell(index, col)
            cell.text = text
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(11)
            para.font.bold = col == 1
            para.font.color.rgb = INK_SECONDARY if col == 0 else INK
            if col == 1:
                para.alignment = PP_ALIGN.RIGHT
    _table(
        slide,
        [
            {"key": "year", "label": "Year", "format": "{:,.0f}"},
            {"key": "revenue", "label": "Revenue", "numeric": True, "format": "{:,.0f}"},
            {"key": "growth", "label": "Growth", "numeric": True, "format": "{:.1%}"},
            {"key": "ebitda", "label": "EBITDA", "numeric": True, "format": "{:,.0f}"},
            {"key": "ebit", "label": "EBIT", "numeric": True, "format": "{:,.0f}"},
            {"key": "taxes", "label": "Taxes", "numeric": True, "format": "{:,.0f}"},
            {"key": "capex", "label": "Capex", "numeric": True, "format": "{:,.0f}"},
            {"key": "fcf", "label": "UFCF", "numeric": True, "format": "{:,.0f}"},
            {"key": "pv_fcf", "label": "PV", "numeric": True, "format": "{:,.0f}"},
        ],
        projections, top=Inches(4.6),
    )
    _note(slide, payload.get("note", ""), top=Inches(6.6))


def _slide_sensitivity(prs, slide_data: dict[str, Any]) -> None:
    slide = _blank(prs)
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    _add_heatmap_table(slide, slide_data.get("payload", {}))
    payload = slide_data.get("payload", {})
    _note(
        slide,
        f"Value per share ({payload.get('currency', 'USD')}); current share price "
        f"{_symbol(payload.get('currency'))}{payload.get('reference_price', 0):,.2f}.",
        top=Inches(5.0),
    )


def _slide_buyers(prs, slide_data: dict[str, Any]) -> None:
    payload = slide_data.get("payload", {})
    slide = _blank(prs)
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    _table(
        slide,
        [
            {"key": "name", "label": "Counterparty"},
            {"key": "type", "label": "Type"},
            {"key": "enterprise_value", "label": "EV (USDm)", "numeric": True, "format": "{:,.0f}"},
            {"key": "relative_size", "label": "EV vs target", "numeric": True, "format": "{:,.2f}x"},
            {"key": "rationale", "label": "Rationale"},
        ],
        payload.get("candidates", []), top=CONTENT_TOP, font_size=9.0,
    )


def _slide_evidence(prs, slide_data: dict[str, Any]) -> None:
    items = slide_data.get("payload", {}).get("items", [])
    slide = _blank(prs)
    _heading(slide, slide_data.get("title", ""), slide_data.get("subtitle"))
    if not items:
        _bullets(slide, ["No evidence pinned. Pin passages from the research module to build this appendix."])
        return
    bullets = []
    for item in items:
        source = " · ".join(
            str(part) for part in (item.get("ticker"), item.get("title"), item.get("published_at")) if part
        )
        bullets.append({
            "text": f"“{item.get('quote', '')}”",
            "source": {
                "title": source, "doc_type": item.get("doc_type", ""),
                "published_at": "", "document_id": item.get("document_id"),
            },
        })
    _bullets(slide, bullets)


WRITERS = {
    "cover": _slide_cover,
    "disclaimer": _slide_text,
    "executive_summary": _slide_executive_summary,
    "company_overview": _slide_company_overview,
    "market_overview": _slide_market_overview,
    "financial_performance": _slide_financials,
    "valuation_summary": _slide_valuation_summary,
    "trading_comps": _slide_trading_comps,
    "precedent_transactions": _slide_precedents,
    "dcf": _slide_dcf,
    "dcf_sensitivity": _slide_sensitivity,
    "buyers": _slide_buyers,
    "process": _slide_text,
    "evidence_appendix": _slide_evidence,
}


def build_presentation(book: dict[str, Any]) -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for slide_data in book.get("slides", []):
        WRITERS.get(slide_data.get("kind", ""), _slide_text)(prs, slide_data)
    return prs


def write_pptx(book: dict[str, Any], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    build_presentation(book).save(str(path))
    return path
